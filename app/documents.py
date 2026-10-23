from io import BytesIO
from zipfile import ZipFile
from docx import Document
from fastapi import HTTPException
import re

MAX_UPLOAD = 5 * 1024 * 1024

def parse_docx(data: bytes) -> str:
    try:
        with ZipFile(BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > 1000 or sum(m.file_size for m in members) > 20*1024*1024:
                raise HTTPException(413, "Document expands beyond the parsing limit")
            if "word/document.xml" not in archive.namelist():
                raise ValueError()
        document = Document(BytesIO(data))
        lines = [p.text for p in document.paragraphs]
        for table in document.tables:
            lines += [" | ".join(cell.text for cell in row.cells) for row in table.rows]
        text = "\n".join(line.strip() for line in lines if line.strip())
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(422, "Cannot parse this file as a DOCX document") from exc
    if not text:
        raise HTTPException(422, "No readable text found; image-only notes are unsupported")
    if len(text) > 100000:
        raise HTTPException(413, "Notes must contain at most 100,000 text characters")
    return text

def retrieve(text, query):
    # Bounded, transparent lexical retrieval for a two-day sprint; no vector database.
    chunks = [text[i:i+1200] for i in range(0, len(text), 1000)]
    terms = set(re.findall(r"\w+", query.lower()))
    ranked = sorted(enumerate(chunks), key=lambda pair: len(terms & set(re.findall(r"\w+", pair[1].lower()))), reverse=True)
    return [{"chunk":i+1, "text":chunk} for i,chunk in ranked[:4]]
