from datetime import date,timedelta
from io import BytesIO
from zipfile import ZipFile
import asyncio
import httpx
import jwt
import pytest
from docx import Document
from fastapi import HTTPException
from app.automation import run_daily
from app.documents import parse_docx,retrieve,MAX_UPLOAD
from app.ai import generate
from app.schemas import QuizRequest,StudyRequest
from app.content import questions

def doc_bytes():
    doc=Document();doc.add_paragraph("Azure Blob Storage holds co-op produce photographs.")
    table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='Elasticity';table.cell(0,1).text='Capacity follows demand'
    output=BytesIO();doc.save(output);return output.getvalue()

def upload(client,auth):
    return client.post('/documents',headers=auth,files={'file':('notes.docx',doc_bytes(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})

def test_auth_and_health(client,auth):
    assert client.get('/health').json()['unofficial'] is True
    assert client.get('/readiness').status_code==401
    assert client.post('/auth/login',json={'email':'learner@example.com','password':'strong-password-123'}).status_code==200
    assert client.post('/auth/login',json={'email':'learner@example.com','password':'incorrect-password'}).status_code==401
    assert client.post('/auth/register',json={'email':'learner@example.com','password':'strong-password-123'}).status_code==409

def test_expired_and_tampered_tokens(client,auth):
    expired=jwt.encode({'sub':'1','exp':1,'iss':'az900-study'},'test-secret-that-is-at-least-32-characters',algorithm='HS256')
    assert client.get('/readiness',headers={'Authorization':f'Bearer {expired}'}).status_code==401
    assert client.get('/readiness',headers={'Authorization':'Bearer nonsense'}).status_code==401

def test_parse_and_retrieve():
    text=parse_docx(doc_bytes());assert 'Elasticity | Capacity follows demand' in text
    assert 'Blob' in retrieve(text,'Blob photos')[0]['text']

def test_upload_isolation_and_deletion(client,auth):
    r=upload(client,auth);assert r.status_code==201;r=r.json()
    assert r['characters']>20
    other=client.post('/auth/register',json={'email':'other@example.com','password':'another-password-123'}).json()
    headers={'Authorization':'Bearer '+other['access_token']}
    assert client.get('/documents',headers=headers).json()==[]
    assert client.post('/buddy',headers=headers,json={'document_id':r['id'],'prompt':'Explain Blob'}).status_code==404
    assert client.delete('/documents/'+str(r['id']),headers=headers).status_code==404
    assert client.delete('/documents/'+str(r['id']),headers=auth).status_code==204

def test_upload_failures(client,auth):
    assert client.post('/documents',headers=auth,files={'file':('notes.txt',b'hello')}).status_code==415
    assert client.post('/documents',headers=auth,files={'file':('notes.docx',b'broken zip')}).status_code==422
    assert client.post('/documents',headers=auth,files={'file':('big.docx',b'x'*(MAX_UPLOAD+1))}).status_code==413
    doc=Document();data=BytesIO();doc.save(data)
    assert client.post('/documents',headers=auth,files={'file':('empty.docx',data.getvalue())}).status_code==422

def test_archive_bomb_rejected():
    data=BytesIO()
    with ZipFile(data,'w',compression=8) as z:z.writestr('word/document.xml',b'x'*(21*1024*1024))
    with pytest.raises(HTTPException) as error:parse_docx(data.getvalue())
    assert error.value.status_code==413

def test_quiz_scoring_and_history(client,auth):
    r=client.post('/quizzes',headers=auth,json={'topic':'services','count':2}).json()
    assert 'answer' not in r['questions'][0]
    hist=client.get('/history',headers=auth).json()[0]
    assert 'answer' not in hist['payload']['questions'][0]
    body={'question_index':0,'selected':0}
    a=client.post(f"/quizzes/{r['id']}/answers",headers=auth,json=body)
    assert a.json()['correct'] and len(a.json()['explanation'])>20
    assert client.post(f"/quizzes/{r['id']}/answers",headers=auth,json=body).status_code==409
    stats=client.get('/readiness',headers=auth).json()
    assert stats['topics'][1]['accuracy']==100 and stats['practice_estimate']==37.5
    assert stats['sufficient_evidence'] is False
    assert client.post(f"/quizzes/{r['id']}/answers",headers=auth,json={'question_index':99,'selected':0}).status_code==422

def test_quiz_ownership(client,auth):
    quiz=client.post('/quizzes',headers=auth,json={}).json()
    other=client.post('/auth/register',json={'email':'outsider@example.com','password':'another-password-123'}).json()
    headers={'Authorization':'Bearer '+other['access_token']}
    assert client.post(f"/quizzes/{quiz['id']}/answers",headers=headers,json={'question_index':0,'selected':0}).status_code==404
    assert client.get('/history',headers=headers).json()==[]

def test_flashcard_spaced_repetition(client,auth):
    card=client.get('/flashcards',headers=auth).json()[0]
    missed=client.post(f"/flashcards/{card['id']}/review",headers=auth,json={'remembered':False}).json()
    assert missed['interval_days']==0 and missed['due']==str(date.today())
    good=client.post(f"/flashcards/{card['id']}/review",headers=auth,json={'remembered':True}).json()
    assert good['interval_days']==1
    assert client.post(f"/flashcards/{card['id']}/review",headers=auth,json={'remembered':True}).status_code==409
    assert len(client.get('/flashcards',headers=auth).json())==11

def test_plan_and_daily_automation(client,auth):
    r=client.post('/plans',headers=auth,json={'exam_date':str(date.today()+timedelta(days=14)),'hours_per_week':7,'weak_topics':['cloud']})
    assert r.status_code==201
    sessions=r.json()['sessions'];assert len(sessions)==14
    assert sum(s['minutes'] for s in sessions)==840
    assert {s['topic'] for s in sessions}=={'cloud','services','governance'}
    assert run_daily()==1 and run_daily()==0
    assert len(client.get('/reminders',headers=auth).json())==1
    assert client.post('/plans',headers=auth,json={'exam_date':str(date.today()),'hours_per_week':7}).status_code==422

def test_validation_and_demo_boundaries(client,auth):
    assert client.post('/quizzes',headers=auth,json={'topic':'invalid'}).status_code==422
    assert client.post('/quizzes',headers=auth,json={'count':9}).status_code==422
    document=upload(client,auth).json()
    assert client.post('/buddy',headers=auth,json={'document_id':document['id']}).status_code==503
    assert client.post('/buddy',headers=auth,json={'language':'Kiswahili'}).status_code==503
    assert 'Offline demonstration' in client.post('/buddy',headers=auth,json={}).json()['answer']

def test_live_generation_contract(monkeypatch):
    monkeypatch.setenv('AI_MODE','live');monkeypatch.setenv('LLM_API_KEY','test-key')
    async def post(self,url,**kwargs):
        assert kwargs['json']['messages'][0]['role']=='system'
        assert 'document_excerpts' in kwargs['json']['messages'][1]['content']
        import json
        return httpx.Response(200,json={'choices':[{'message':{'content':json.dumps({'questions':questions('cloud',1)})}}]},request=httpx.Request('POST',url))
    monkeypatch.setattr(httpx.AsyncClient,'post',post)
    r=asyncio.run(generate('quiz',QuizRequest(count=1),[{'chunk':1,'text':'study data'}]))
    assert r['mode']=='live' and len(r['questions'])==1

def test_live_failure_is_explicit(monkeypatch):
    monkeypatch.setenv('AI_MODE','live');monkeypatch.setenv('LLM_API_KEY','test-key')
    async def post(self,url,**kwargs):return httpx.Response(200,json={'choices':[{'message':{'content':'not-json'}}]},request=httpx.Request('POST',url))
    monkeypatch.setattr(httpx.AsyncClient,'post',post)
    with pytest.raises(HTTPException) as err:asyncio.run(generate('quiz',QuizRequest(),[]))
    assert err.value.status_code==502

def test_mcp_tool_uses_bearer(monkeypatch):
    from app.mcp_server import get_my_readiness
    monkeypatch.setenv('STUDY_ACCESS_TOKEN','learner-token')
    async def get(self,url,**kwargs):
        assert kwargs['headers']['Authorization']=='Bearer learner-token'
        return httpx.Response(200,json={'practice_estimate':42},request=httpx.Request('GET',url))
    monkeypatch.setattr(httpx.AsyncClient,'get',get)
    assert asyncio.run(get_my_readiness())['practice_estimate']==42
